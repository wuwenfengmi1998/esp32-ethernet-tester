#!/usr/bin/env python3
"""Generate a sales slick (one-pager) for the ESP32 Ethernet Tester."""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color,
        qn('w:val'): 'clear',
        qn('w:color'): 'auto'
    })
    shading.append(shd)

def main():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('ESP32 Ethernet Tester')
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0, 70, 127)

    # Subtitle
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run('Portable Network Validation & Security Assessment Platform')
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph()

    # Tagline
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('A pocket-sized bench tool for stress-testing, validating, and security-auditing '
                    'Ethernet switch ports. Full raw-frame control, protocol simulation, and '
                    'offensive testing in a single PoE-powered device.')
    run.font.size = Pt(11)
    run.italic = True

    doc.add_paragraph()

    # --- Feature Categories ---
    categories = [
        ('Layer 1/2 Testing', [
            'Raw Ethernet frame TX/RX (W5500 MACRAW mode)',
            'Error injection: giant, jumbo, broadcast, multicast, PAUSE, pattern storms',
            'Continuous stress generation at configurable rates',
            'PHY speed/duplex control (auto, 100FD, 100HD, 10FD, 10HD)',
            'RFC 2544 benchmarks: throughput, latency, frame-loss, back-to-back',
        ]),
        ('Network Discovery & Reconnaissance', [
            'LLDP / CDP passive decode and active advertisement',
            'Passive host/protocol mapping',
            'ICMP ping sweep and traceroute',
            'ARP scan (L2 host discovery)',
            'TCP SYN port scan with banner grab (30+ service probes)',
            'SNMP community probe, IP sweep, and write-access test',
            'Combined network assessment with CIDR support',
        ]),
        ('Protocol Testing', [
            'Full DHCP DORA verification + 6 failure/abuse scenarios',
            'DHCPv6 stateful server discovery (SOLICIT)',
            'DNS A-record resolution and rogue DNS responder',
            'mDNS .local resolution and reachability probe',
            'IPv6 NDP decode (RS/RA/NS/NA)',
            'FHRP (HSRP/VRRP) passive decode',
        ]),
        ('Security Assessment', [
            '802.1X supplicant: EAP-MD5, EAP-TLS, PEAPv0/MSCHAPv2, EAP-TTLS',
            'Rogue authenticator with credential harvesting (hashcat output)',
            'EAPOL-Start flood, spoofed logoff, MAB probe',
            'L2 attacks: VLAN inject, Q-in-Q hop, DTP spoof, CAM flood, STP root claim',
            'ARP MITM, gratuitous ARP, ARP storm',
            'Rogue Router Advertisement (SLAAC takeover)',
            'HSRP/VRRP gateway hijack',
            'Rogue DHCP/DHCPv6 server, DNS spoofing',
            'MAC randomization per-operation and per-host',
        ]),
        ('Management & Automation', [
            'Wi-Fi web UI: 8-tab interface with real-time output streaming',
            'Serial CLI with 100+ commands',
            'Scripting engine: loops, conditionals, delays, cron scheduling',
            'PCAP capture to SD card (downloadable via web)',
            'Output logging with timestamped log files',
            'OTA firmware updates (TFTP + HTTP upload)',
            'WireGuard VPN tunnel for secure remote management',
            'HTTPS with basic authentication',
            'NVS persistence for all configuration',
        ]),
    ]

    for cat_name, features in categories:
        # Category heading
        h = doc.add_heading(cat_name, level=2)
        h.runs[0].font.color.rgb = RGBColor(0, 70, 127)

        for feat in features:
            p = doc.add_paragraph(feat, style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)

    doc.add_paragraph()

    # --- Hardware Specs ---
    doc.add_heading('Hardware Platform', level=2).runs[0].font.color.rgb = RGBColor(0, 70, 127)

    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    specs = [
        ('MCU', 'ESP32-S3 (dual-core 240 MHz, 16 MB flash, 8 MB PSRAM)'),
        ('Ethernet', 'WIZnet W5500 (onboard, PoE-capable)'),
        ('Form Factor', 'Waveshare ESP32-S3-POE-ETH (compact single-board)'),
        ('Power', 'USB-C or 802.3af PoE'),
        ('Storage', 'Micro-SD card (FAT32) + LittleFS internal'),
    ]
    for i, (label, value) in enumerate(specs):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value
        # Bold the label column
        table.cell(i, 0).paragraphs[0].runs[0].bold = True
        set_cell_shading(table.cell(i, 0), 'E8F0F8')

    doc.add_paragraph()

    # --- Key Differentiators ---
    doc.add_heading('Key Differentiators', level=2).runs[0].font.color.rgb = RGBColor(0, 70, 127)

    diffs = [
        'True raw-frame control -- crafts any Ethernet frame at the MAC layer',
        'All-in-one: L1 stress testing through L7 service identification in one device',
        'No laptop required -- standalone operation via Wi-Fi web UI or scripted automation',
        'Pocket-sized and PoE-powered -- deploy anywhere with a single cable',
        'Arm/disarm safety model -- offensive tests locked behind explicit authorization',
        'Open firmware with OTA updates -- continuously expanding capability set',
    ]
    for d in diffs:
        p = doc.add_paragraph(d, style='List Bullet')
        p.paragraph_format.space_after = Pt(3)

    doc.add_paragraph()

    # --- Use Cases ---
    doc.add_heading('Use Cases', level=2).runs[0].font.color.rgb = RGBColor(0, 70, 127)

    cases = [
        ('Switch Port Validation', 'Verify port security, storm control, BPDU guard, DHCP snooping, and 802.1X enforcement on new deployments.'),
        ('Network Penetration Testing', 'Assess L2/L3 attack surface: VLAN hopping, ARP poisoning, rogue services, credential harvesting.'),
        ('Troubleshooting', 'LLDP/CDP neighbor discovery, DHCP diagnostics, DNS resolution testing, link-layer analysis.'),
        ('Performance Benchmarking', 'RFC 2544 throughput/latency/loss measurements for acceptance testing.'),
        ('Compliance Auditing', 'Verify 802.1X enforcement, SNMP hardening, and protocol filtering policies.'),
    ]
    for title_text, desc in cases:
        p = doc.add_paragraph()
        run = p.add_run(title_text + ' -- ')
        run.bold = True
        p.add_run(desc)
        p.paragraph_format.space_after = Pt(4)

    output = 'ESP32_Ethernet_Tester_Sales_Slick.docx'
    doc.save(output)
    print(f'Saved: {output}')

if __name__ == '__main__':
    main()
