#!/usr/bin/env python3
"""Convert README.md to a Word document."""

import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def parse_table(lines):
    """Parse markdown table lines into rows of cells."""
    rows = []
    for line in lines:
        line = line.strip().strip('|')
        cells = [c.strip() for c in line.split('|')]
        rows.append(cells)
    # Remove separator row (---|----|---)
    if len(rows) > 1 and all(re.match(r'^[-:]+$', c) for c in rows[1]):
        rows.pop(1)
    return rows

def add_table(doc, rows):
    """Add a formatted table to the document."""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    # Normalize all rows to same column count
    for r in rows:
        while len(r) < ncols:
            r.append('')
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = row[j] if j < len(row) else ''
            cell = table.cell(i, j)
            # Strip bold markers
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', cell_text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            cell.text = text
            # Bold header row
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()

def add_code_block(doc, code_lines):
    """Add a code block as a formatted paragraph."""
    p = doc.add_paragraph()
    p.style = doc.styles['No Spacing']
    for i, line in enumerate(code_lines):
        run = p.add_run(line)
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        if i < len(code_lines) - 1:
            p.add_run('\n')

def process_inline(paragraph, text):
    """Process inline markdown (bold, code, links) and add runs."""
    # Pattern for bold, code, and links
    pattern = r'(\*\*(.+?)\*\*|`(.+?)`|\[(.+?)\]\((.+?)\))'
    last_end = 0
    for m in re.finditer(pattern, text):
        # Add text before match
        if m.start() > last_end:
            paragraph.add_run(text[last_end:m.start()])
        if m.group(2):  # bold
            run = paragraph.add_run(m.group(2))
            run.bold = True
        elif m.group(3):  # code
            run = paragraph.add_run(m.group(3))
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        elif m.group(4):  # link
            paragraph.add_run(m.group(4))
        last_end = m.end()
    if last_end < len(text):
        paragraph.add_run(text[last_end:])

def main():
    with open('README.md', 'r') as f:
        lines = f.readlines()

    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            # Strip inline formatting for heading
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # Code block
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            add_code_block(doc, code_lines)
            i += 1  # skip closing ```
            continue

        # Table
        if '|' in line and i + 1 < len(lines) and '|' in lines[i + 1]:
            table_lines = []
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                table_lines.append(lines[i].rstrip('\n'))
                i += 1
            rows = parse_table(table_lines)
            add_table(doc, rows)
            continue

        # Blockquote
        if line.startswith('>'):
            quote_text = line.lstrip('> ').strip()
            i += 1
            while i < len(lines) and lines[i].startswith('>'):
                quote_text += ' ' + lines[i].lstrip('> ').strip()
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.style = doc.styles['No Spacing']
            process_inline(p, quote_text)
            continue

        # Bullet list
        if re.match(r'^[-*]\s', line):
            while i < len(lines) and re.match(r'^[-*]\s', lines[i]):
                text = re.sub(r'^[-*]\s+', '', lines[i].rstrip('\n'))
                p = doc.add_paragraph(style='List Bullet')
                process_inline(p, text)
                i += 1
            continue

        # Numbered list
        if re.match(r'^\d+\.\s', line):
            while i < len(lines) and re.match(r'^\d+\.\s', lines[i]):
                text = re.sub(r'^\d+\.\s+', '', lines[i].rstrip('\n'))
                p = doc.add_paragraph(style='List Number')
                process_inline(p, text)
                i += 1
            continue

        # Regular paragraph
        para_text = line
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith('#') \
                and not lines[i].startswith('```') and not lines[i].startswith('>') \
                and not lines[i].strip().startswith('|') and not lines[i].strip() == '---' \
                and not re.match(r'^[-*]\s', lines[i]) and not re.match(r'^\d+\.\s', lines[i]):
            para_text += ' ' + lines[i].strip()
            i += 1
        p = doc.add_paragraph()
        process_inline(p, para_text)

    output = 'README.docx'
    doc.save(output)
    print(f'Saved: {output}')

if __name__ == '__main__':
    main()
